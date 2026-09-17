"""Build reproducible Excel deliverables from Meteora's public DLMM Data API.

Run: python build_project.py
The script stores the exact API payloads next to the workbooks so that the
portfolio claims remain auditable rather than using invented chain metrics.
"""
from __future__ import annotations

import json
import shutil
import time
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

import requests
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
API = "https://dlmm.datapi.meteora.ag"
SOURCE_CV = ROOT.parent / "tech" / "CV_tech_ai.docx"
OUT_CV = ROOT / "Yongyi_YE_Web3_Data_BI_Resume.docx"
OUT_REPORT = ROOT / "03_Meteora_DLMM_Operations_Analysis_Report.docx"
NAVY = "17365D"
TEAL = "00A6A6"
PALE = "EAF4F4"
THIN = Side(style="thin", color="D9E2F3")


def get_json(path: str):
    """Fetch public data with retries, keeping the build resilient to transient TLS errors."""
    last = None
    for attempt in range(4):
        try:
            response = requests.get(API + path, timeout=45)
            response.raise_for_status()
            return response.json()
        except requests.RequestException as exc:
            last = exc
            time.sleep(2 * (attempt + 1))
    raise RuntimeError(f"Could not retrieve {path}: {last}")


def style_header(ws, row, end_col):
    for cell in ws[row][0:end_col]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=THIN)


def fit(ws):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        width = min(max(len(str(c.value or "")) for c in col) + 2, 42)
        ws.column_dimensions[letter].width = max(width, 12)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def pool_rows(pools):
    rows = []
    for p in pools:
        volume = p.get("volume", {})
        fees = p.get("fees", {})
        txns = p.get("transactions", {})
        tvl = float(p.get("tvl") or 0)
        vol24 = float(volume.get("24h") or 0)
        fee24 = float(fees.get("24h") or 0)
        rows.append({
            "pool": p.get("name", "Unlabelled pool"),
            "address": p.get("address", ""),
            "token_x": p.get("token_x", {}).get("symbol", ""),
            "token_y": p.get("token_y", {}).get("symbol", ""),
            "created": datetime.fromtimestamp((p.get("created_at") or 0) / 1000, tz=timezone.utc).date(),
            "tvl": tvl,
            "vol24": vol24,
            "fees24": fee24,
            "turnover": vol24 / tvl if tvl else 0,
            "fee_yield": fee24 / tvl if tvl else 0,
            "apr": float(p.get("apr") or 0),
            "apy": float(p.get("apy") or 0),
            "tx24": float(txns.get("24h") or 0),
            "farm": bool(p.get("has_farm")),
        })
    return sorted(rows, key=lambda r: r["vol24"], reverse=True)


def build_raw_workbook(rows, pulled_at):
    wb = Workbook()
    ws = wb.active
    ws.title = "Pool Snapshot"
    columns = ["Pool", "Pool Address", "Token X", "Token Y", "Created (UTC)", "TVL USD", "24h Volume USD", "24h Fees USD", "Volume / TVL", "Fee Yield (24h)", "APR", "APY", "24h Transactions", "Farm Enabled"]
    ws.append(columns)
    for r in rows:
        ws.append([r["pool"], r["address"], r["token_x"], r["token_y"], r["created"], r["tvl"], r["vol24"], r["fees24"], r["turnover"], r["fee_yield"], r["apr"], r["apy"], r["tx24"], r["farm"]])
    style_header(ws, 1, len(columns))
    for row in range(2, ws.max_row + 1):
        for col in (6, 7, 8): ws.cell(row, col).number_format = '$#,##0.00'
        for col in (9, 10, 11, 12): ws.cell(row, col).number_format = '0.00%'
        ws.cell(row, 13).number_format = '#,##0'
    ws.conditional_formatting.add(f"I2:I{ws.max_row}", ColorScaleRule(start_type='min', start_color='FDE9D9', mid_type='percentile', mid_value=50, mid_color='FFF2CC', end_type='max', end_color='D9EAD3'))
    fit(ws)

    meta = wb.create_sheet("Data Notes")
    meta.append(["Field", "Definition / source"])
    meta.append(["Source", "Meteora DLMM Data API /pools (public, read-only)"])
    meta.append(["Snapshot time (UTC)", pulled_at.isoformat()])
    meta.append(["Scope", "First API page of currently indexed DLMM pools; it is a reproducible snapshot, not a claim about all protocol pools."])
    meta.append(["TVL", "USD value of liquidity shown by the API at extraction."])
    meta.append(["Volume / TVL", "24h trading volume divided by current TVL; a liquidity-utilisation proxy, not a return."])
    meta.append(["Fee Yield (24h)", "24h fees divided by current TVL; annualisation is intentionally excluded."])
    style_header(meta, 1, 2); fit(meta)
    wb.save(ROOT / "01_Raw_Pool_Snapshot.xlsx")


def build_analysis_workbook(rows, trend, pulled_at):
    wb = Workbook()
    dashboard = wb.active
    dashboard.title = "Dashboard"
    dashboard.merge_cells("A1:F1")
    dashboard["A1"] = "Meteora DLMM — Liquidity Pool Operations Dashboard"
    dashboard["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    dashboard["A1"].fill = PatternFill("solid", fgColor=NAVY)
    dashboard["A1"].alignment = Alignment(horizontal="center")
    dashboard.append([])
    dashboard.append(["Snapshot (UTC)", pulled_at.strftime("%Y-%m-%d %H:%M"), "Pools in sample", len(rows), "Data source", "Meteora public API"])
    dashboard.append(["Sample TVL", "=SUM('Pool Benchmark'!F2:F11)", "24h Volume", "=SUM('Pool Benchmark'!G2:G11)", "24h Fees", "=SUM('Pool Benchmark'!H2:H11)"])
    dashboard.append(["Weighted turnover", "=D4/B4", "Weighted fee yield", "=F4/B4", "Top pool by 24h volume", "='Pool Benchmark'!A2"])
    for c in ("B4", "D4", "F4"): dashboard[c].number_format = '$#,##0.00'
    for c in ("B5", "D5"): dashboard[c].number_format = '0.00%'
    for cell in dashboard[3]: cell.fill = PatternFill("solid", fgColor=PALE); cell.font = Font(bold=True)
    for cell in dashboard[4]: cell.font = Font(bold=True, color=NAVY)
    for cell in dashboard[5]: cell.font = Font(bold=True, color=NAVY)
    dashboard["A8"] = "Analyst questions answered"
    dashboard["A8"].font = Font(bold=True, color="FFFFFF")
    dashboard["A8"].fill = PatternFill("solid", fgColor=TEAL)
    notes = [
        "1. Which sampled pools produce the most trading activity per unit of liquidity? (Volume / TVL)",
        "2. Where do fee monetisation and liquidity utilisation diverge? (Fee yield vs. turnover)",
        "3. Is protocol-level daily volume rising, stable, or volatile over the latest available observations?",
        "Use this workbook as an operating diagnostic—not investment advice or a performance forecast.",
    ]
    for note in notes: dashboard.append([note])
    dashboard.column_dimensions["A"].width = 32
    for c in "BCDEF": dashboard.column_dimensions[c].width = 19

    pools = wb.create_sheet("Pool Benchmark")
    headers = ["Pool", "Pool Address", "Token X", "Token Y", "Created (UTC)", "TVL USD", "24h Volume USD", "24h Fees USD", "Volume / TVL", "Fee Yield (24h)", "APR", "APY", "24h Txns", "Farm"]
    pools.append(headers)
    for r in rows[:10]:
        pools.append([r["pool"], r["address"], r["token_x"], r["token_y"], r["created"], r["tvl"], r["vol24"], r["fees24"], r["turnover"], r["fee_yield"], r["apr"], r["apy"], r["tx24"], r["farm"]])
    style_header(pools, 1, len(headers))
    for row in range(2, pools.max_row + 1):
        for col in (6,7,8): pools.cell(row, col).number_format = '$#,##0.00'
        for col in (9,10,11,12): pools.cell(row, col).number_format = '0.00%'
        pools.cell(row, 13).number_format = '#,##0'
    pools.conditional_formatting.add(f"I2:I{pools.max_row}", ColorScaleRule(start_type='min', start_color='F4CCCC', mid_type='percentile', mid_value=50, mid_color='FFF2CC', end_type='max', end_color='D9EAD3'))
    fit(pools)
    bar = BarChart()
    bar.title = "24h Volume by Pool"
    bar.y_axis.title = "USD"
    bar.add_data(Reference(pools, min_col=7, min_row=1, max_row=pools.max_row), titles_from_data=True)
    bar.set_categories(Reference(pools, min_col=1, min_row=2, max_row=pools.max_row))
    bar.height = 8; bar.width = 15
    dashboard.add_chart(bar, "A14")

    trend_sheet = wb.create_sheet("Protocol Daily Trend")
    trend_sheet.append(["Date (UTC)", "Daily Volume USD", "7D Moving Average", "Day-over-day change"])
    for i, point in enumerate(trend):
        date = datetime.fromtimestamp(point["timestamp"], tz=timezone.utc).date()
        row = i + 2
        trend_sheet.append([date, point["clean"], f"=AVERAGE(B{max(2, row-6)}:B{row})", f"=IFERROR(B{row}/B{row-1}-1,\"\")" if row > 2 else ""])
    style_header(trend_sheet, 1, 4)
    for row in range(2, trend_sheet.max_row + 1):
        trend_sheet.cell(row, 2).number_format = '$#,##0.00'
        trend_sheet.cell(row, 3).number_format = '$#,##0.00'
        trend_sheet.cell(row, 4).number_format = '0.00%'
    fit(trend_sheet)
    line = LineChart()
    line.title = "Protocol Daily Volume & 7D Average"
    line.y_axis.title = "USD"
    line.add_data(Reference(trend_sheet, min_col=2, max_col=3, min_row=1, max_row=trend_sheet.max_row), titles_from_data=True)
    line.set_categories(Reference(trend_sheet, min_col=1, min_row=2, max_row=trend_sheet.max_row))
    line.height = 8; line.width = 15
    dashboard.add_chart(line, "N14")

    definitions = wb.create_sheet("Metric Guide")
    definitions.append(["Metric", "Business interpretation", "Caveat"])
    definitions.append(["TVL", "Current liquidity capacity in a pool.", "Point-in-time and price-sensitive."])
    definitions.append(["24h Volume", "Recent demand / trading activity flowing through a pool.", "May spike around launches or market events."])
    definitions.append(["Volume / TVL", "Liquidity utilisation proxy for allocating monitoring attention.", "Not a profitability or risk measure by itself."])
    definitions.append(["Fee Yield (24h)", "Daily gross fees relative to TVL; useful for comparing monetisation intensity.", "No annualisation; fees can be highly volatile."])
    definitions.append(["APR / APY", "API-displayed pool return indicators.", "Do not treat as guaranteed returns; display values may include assumptions."])
    definitions.append(["Source and refresh", "Meteora DLMM Data API: /pools and /stats/daily/volume.", "Re-run build_project.py to capture a new dated snapshot."])
    style_header(definitions, 1, 3); fit(definitions)
    wb.save(ROOT / "02_Meteora_DLMM_Operations_Analysis.xlsx")


def patch_resume():
    """Clone the latest tech CV and replace the least role-relevant retail project."""
    shutil.copy2(SOURCE_CV, OUT_CV)
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    replacements = {
        "Interest-based E-commerce User Behavior & Sales Analysis | Data Analyst[Dashboard Link]Aug 2025 – Oct 2025": "Meteora DLMM Liquidity Pool Operations Analysis | Independent Data Analytics Project [Portfolio Link] Sep 2026",
        "Decision & Temporal Strategy: Identified the \"Golden 120s\" impulse buying window and 20:00-23:00 sales peak, guiding UX optimization and hourly budget allocation.": "On-chain Operations: Collected reproducible TVL, volume, fee, transaction and yield snapshots from Meteora's public DLMM API; created a metric guide and documented the refresh logic.",
        "Ad Efficiency: Detected a significant ROI gap between top slots (e.g., Vprol56) and long-tail slots; proposed cutting the bottom 20% budget to optimize ROAS.": "Liquidity & Growth Diagnosis: Benchmarked pools with Volume/TVL and 24-hour Fee Yield metrics; surfaced product questions around liquidity allocation, demand volatility and LP experience through an Excel BI dashboard.",
    }
    with zipfile.ZipFile(OUT_CV, "r") as zin:
        contents = {name: zin.read(name) for name in zin.namelist()}
    root = ET.fromstring(contents["word/document.xml"])
    found = set()
    for paragraph in root.iter(ns + "p"):
        nodes = list(paragraph.iter(ns + "t"))
        text = "".join(node.text or "" for node in nodes)
        for old, new in replacements.items():
            if old in text:
                nodes[0].text = text.replace(old, new)
                for node in nodes[1:]: node.text = ""
                found.add(old)
                break
    if len(found) != len(replacements):
        missing = set(replacements) - found
        raise RuntimeError(f"Resume text was not safely updated; missing paragraphs: {missing}")
    contents["word/document.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with zipfile.ZipFile(OUT_CV, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, payload in contents.items(): zout.writestr(name, payload)


def build_report_docx(rows, trend, pulled_at):
    """Create an image-ready portfolio report. Tableau figures are intentionally marked as placeholders."""
    total_tvl = sum(r["tvl"] for r in rows)
    total_volume = sum(r["vol24"] for r in rows)
    total_fees = sum(r["fees24"] for r in rows)
    top = rows[0]
    top_share = top["vol24"] / total_volume if total_volume else 0
    trend_values = [p["clean"] for p in trend]
    latest_7 = sum(trend_values[-7:]) / min(7, len(trend_values)) if trend_values else 0
    previous_7 = sum(trend_values[-14:-7]) / 7 if len(trend_values) >= 14 else 0
    movement = (latest_7 / previous_7 - 1) if previous_7 else None

    sections = [
        ("title", "Meteora DLMM Liquidity Pool Operations Analysis"),
        ("subtitle", "Web3 Business Intelligence Portfolio Report | Solana | Public API snapshot"),
        ("h1", "1. Executive Summary"),
        ("body", f"This project treats decentralised liquidity pools as an operating product rather than a trading narrative. It combines a dated public API snapshot of {len(rows)} Meteora DLMM pools with protocol-level daily volume history to identify where liquidity capacity, trading demand and fee monetisation diverge."),
        ("body", f"At extraction ({pulled_at.strftime('%d %b %Y, %H:%M UTC')}), the sampled pools represented ${total_tvl:,.0f} TVL, ${total_volume:,.0f} in 24-hour volume and ${total_fees:,.0f} in 24-hour fees. {top['pool']} was the largest sampled pool by 24-hour volume, accounting for {top_share:.1%} of sample volume. These are point-in-time diagnostic figures, not investment recommendations."),
        ("h1", "2. Business Context & Objective"),
        ("body", "For a decentralised exchange liquidity product, operators need to understand whether liquidity is being used efficiently, whether fee generation follows demand, and whether protocol activity is persistent or event-driven. The analysis answers three questions: (1) which pools show high volume relative to liquidity supplied; (2) where does gross fee intensity differ from utilisation; and (3) how volatile is daily protocol volume?"),
        ("h1", "3. Data & Methodology"),
        ("body", "Data was retrieved from Meteora's public read-only DLMM API endpoints /pools and /stats/daily/volume. The pool snapshot contains TVL, 24-hour volume, fees, transactions, displayed APR/APY and token-pair attributes. Protocol history is filtered to the latest 60 positive daily observations. The included JSON responses preserve the exact source payloads; the Excel workbooks provide the transformation logic."),
        ("body", "Core metrics: Volume / TVL = 24-hour volume ÷ current TVL (liquidity-utilisation proxy). Fee Yield (24h) = 24-hour fees ÷ current TVL (gross daily fee intensity). A 7-day moving average is used to reduce single-day volume noise. Neither metric is a return forecast."),
        ("figure", "FIGURE 1 — Tableau: Pool benchmark scatter plot\nPlace TVL on Columns, 24h Volume on Rows, colour by Fee Yield (24h), size by 24h Transactions, and label by Pool."),
        ("h1", "4. Findings & Operating Implications"),
        ("body", "Pool segmentation: A high Volume / TVL ratio can indicate strong liquidity utilisation and merits monitoring for depth, routing quality and LP retention. High TVL with low utilisation suggests a pool where operators should investigate incentives, token demand or capital efficiency. Fee yield should be assessed alongside utilisation because fee tiers and trading composition can create divergence."),
        ("body", f"Protocol trend: the latest available 7-day average daily volume is ${latest_7:,.0f}. " + (f"This is {movement:+.1%} versus the preceding 7-day average, so the portfolio view should distinguish persistent demand from a short-term event." if movement is not None else "There are not enough prior observations in the selected slice for a prior-week comparison.")),
        ("figure", "FIGURE 2 — Tableau: Daily protocol volume with 7-day moving average\nUse a dual-line time series. Add an annotation for any unusually large daily spike after validating a product or market event."),
        ("h1", "5. Recommendations"),
        ("body", "1. Build a weekly pool-health review prioritising high-volume pools with falling TVL or unusual transaction patterns. 2. Segment pools into utilisation / monetisation quadrants before changing incentives. 3. Pair volume spikes with qualitative event logs—new token launches, campaigns, routing changes or market moves—before deciding that growth is durable. 4. Extend the pipeline with wallet cohorts or position-level data if a business stakeholder needs LP retention analysis."),
        ("h1", "6. Limitations & Reproducibility"),
        ("body", "The pool table is a reproducible API-page snapshot rather than a full audited protocol census. TVL and displayed APR/APY are source-provided, time-sensitive values. The analysis does not make valuation, performance or investment claims. Re-run build_project.py before presenting figures as current; Tableau extracts should be refreshed from the newly generated Excel workbook."),
        ("h1", "Appendix — Deliverables"),
        ("body", "01_Raw_Pool_Snapshot.xlsx — source snapshot and field definitions.\n02_Meteora_DLMM_Operations_Analysis.xlsx — dashboard, benchmark, protocol trend and metric guide.\nsource_pool_snapshot.json and source_protocol_daily_volume.json — API payload archive.\nbuild_project.py — reproducible refresh script."),
    ]
    def para(text, style):
        safe = escape(text).replace("\n", "</w:t><w:br/><w:t>")
        if style == "title": size, bold, color, align, after = 32, True, NAVY, "center", 240
        elif style == "subtitle": size, bold, color, align, after = 18, False, "5B6573", "center", 360
        elif style == "h1": size, bold, color, align, after = 18, True, NAVY, "left", 120
        elif style == "figure": size, bold, color, align, after = 11, True, TEAL, "left", 280
        else: size, bold, color, align, after = 10, False, "222222", "both", 120
        b = '<w:b/>' if bold else ''
        return f'<w:p><w:pPr><w:jc w:val="{align}"/><w:spacing w:after="{after}"/></w:pPr><w:r><w:rPr>{b}<w:color w:val="{color}"/><w:sz w:val="{size*2}"/></w:rPr><w:t xml:space="preserve">{safe}</w:t></w:r></w:p>'
    body = ''.join(para(text, style) for style, text in sections)
    doc = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>{body}<w:sectPr><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1134" w:right="1134" w:bottom="1134" w:left="1134"/></w:sectPr></w:body></w:document>'''.encode("utf-8")
    with zipfile.ZipFile(SOURCE_CV, "r") as zin:
        contents = {name: zin.read(name) for name in zin.namelist()}
    contents["word/document.xml"] = doc
    with zipfile.ZipFile(OUT_REPORT, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, payload in contents.items(): zout.writestr(name, payload)


def main():
    pulled_at = datetime.now(timezone.utc)
    pool_payload = get_json("/pools?limit=10&page=1")
    volume_payload = get_json("/stats/daily/volume")
    (ROOT / "source_pool_snapshot.json").write_text(json.dumps(pool_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    (ROOT / "source_protocol_daily_volume.json").write_text(json.dumps(volume_payload, ensure_ascii=False, indent=2), encoding="utf-8")
    rows = pool_rows(pool_payload["data"])
    points = [p for p in volume_payload["data_points"] if p.get("clean") is not None and p.get("clean", 0) > 0]
    points = sorted(points, key=lambda x: x["timestamp"])[-60:]
    build_raw_workbook(rows, pulled_at)
    build_analysis_workbook(rows, points, pulled_at)
    patch_resume()
    build_report_docx(rows, points, pulled_at)
    print(f"Created 2 Excel workbooks, source JSON, {OUT_CV.name}, and {OUT_REPORT.name} in {ROOT}")


if __name__ == "__main__":
    main()
